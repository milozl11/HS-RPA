"""Read materials + multilingual descriptions from Excel file."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string


def col_to_idx(col: str | int) -> int:
    """Convert 'B' -> 1 (0-based) or accept int already 0-based."""
    if isinstance(col, int):
        return col
    return column_index_from_string(col) - 1


def read_materials(
    path: str | Path,
    material_column: str | int = "B",
    languages: list[dict] | None = None,
    header_row: int = 1,
    sheet: str | None = None,
    # backwards compat
    description_column: str | int | None = None,
) -> list[dict]:
    """Return list of {row, material, descriptions: {CODE: text, ...}, description}.

    `languages` is a list of {code, column, enabled?}. Only enabled languages
    are read. A row is kept only if material is non-empty AND at least one
    description for the requested languages is non-empty.
    """
    if languages is None:
        languages = [
            {"code": "DE", "column": description_column or "J", "enabled": True}
        ]

    active = [
        lng
        for lng in languages
        if lng.get("enabled", True) and lng.get("column") and lng.get("code")
    ]
    if not active:
        return []

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    mat_idx = col_to_idx(material_column)
    lang_idx = [(lng["code"].upper(), col_to_idx(lng["column"])) for lng in active]
    max_idx = max([mat_idx] + [i for _, i in lang_idx])

    items: list[dict] = []
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if i <= header_row:
            continue
        if row is None or len(row) <= max_idx:
            continue
        material = row[mat_idx]
        if material is None:
            continue
        material_s = str(material).strip()
        if not material_s:
            continue

        descriptions: dict[str, str] = {}
        for code, idx in lang_idx:
            val = row[idx]
            if val is None:
                continue
            text = str(val).strip()
            if text:
                descriptions[code] = text

        if not descriptions:
            continue

        items.append(
            {
                "row": i,
                "material": material_s,
                "descriptions": descriptions,
                "description": next(iter(descriptions.values())),
            }
        )
    wb.close()
    return items


if __name__ == "__main__":
    import json
    import sys

    items = read_materials(
        sys.argv[1],
        languages=[
            {"code": "DE", "column": "J", "enabled": True},
            {"code": "EN", "column": "K", "enabled": True},
        ],
    )
    print(json.dumps(items[:5], indent=2, ensure_ascii=False))
    print(f"Total: {len(items)} items")
