"""Reference data for HS/tariff classification.

Maps a SAP GTS numbering scheme (e.g. EDCHSCDEEX) to the tariff column of the
customer reference workbook, then resolves TECDOC -> tariff code.
"""

from __future__ import annotations

import math
import re
from decimal import Decimal, InvalidOperation
from itertools import chain
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

_MISSING_MARKERS = {
    "-",
    "--",
    "N/A",
    "NA",
    "#N/A",
    "NONE",
    "NULL",
    "NOT AVAILABLE",
    "NOT ASSIGNED",
}

# SAP numbering scheme -> (human label, reference workbook column header,
# SAP selection-screen variant). Only the schemes driven by the RPA are kept,
# and the variant already sets the numbering scheme inside SAP.
SCHEMES: dict[str, dict[str, str]] = {
    "EDCHSCAUEX": {
        "label": "Australia",
        "column": "Australia",
        "variant": "STANDARD AU",
    },
    "EDCHSCCNXN": {
        "label": "China new",
        "column": "CHINA",
        "variant": "STANDARD CN",
    },
    "EDCHSCDEEX": {
        "label": "EU",
        "column": "EU TARIC",
        "variant": "STANDARD EU",
    },
    "FHSCUKIM": {
        "label": "Great Britain",
        "column": "Great Britain",
        "variant": "STANDARD GB",
    },
    "EDCHSCINXX": {
        "label": "India",
        "column": "INDIA",
        "variant": "STANDARD IN",
    },
    "EDCHSCNZXX": {
        "label": "New Zealand",
        "column": "New Zealand",
        "variant": "STANDARD NZ",
    },
    "EDCHSCSGXX": {
        "label": "Singapore",
        "column": "SINGAPORE",
        "variant": "STANDARD SG",
    },
    "EDCHSCUSIM": {
        "label": "USA HTS",
        "column": "USA HTS",
        "variant": "STANDARD US",
    },
    "EDCHSCZAXX": {
        "label": "South Africa",
        "column": "South Africa",
        "variant": "STANDARD ZA",
    },
}

# Queue order used when the UI presets every RPA scheme.
RPA_SCHEME_ORDER: list[str] = [
    "EDCHSCDEEX",
    "FHSCUKIM",
    "EDCHSCUSIM",
    "EDCHSCCNXN",
    "EDCHSCINXX",
    "EDCHSCSGXX",
    "EDCHSCAUEX",
    "EDCHSCNZXX",
    "EDCHSCZAXX",
]


def list_schemes() -> list[dict]:
    return [
        {
            "code": code,
            "label": meta["label"],
            "column": meta["column"],
            "variant": meta["variant"],
        }
        for code, meta in sorted(SCHEMES.items())
    ]


def scheme_info(scheme: str) -> dict:
    key = (scheme or "").strip().upper()
    if key not in SCHEMES:
        raise KeyError(f"Schema de numerotare necunoscuta: {scheme!r}")
    return {"code": key, **SCHEMES[key]}


def _header_key(value) -> str:
    """Canonical Excel header used for tolerant, unambiguous matching."""
    text = "" if value is None else str(value)
    text = text.replace("\ufeff", " ").replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip().casefold()


def _plain_number(value) -> str:
    """Render Excel numerics without scientific notation or trailing .0."""
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
        return format(Decimal(str(value)), "f")

    text = str(value).strip()
    if not text:
        return ""
    if re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)[eE][+-]?\d+", text):
        try:
            rendered = format(Decimal(text), "f")
            if "." in rendered:
                rendered = rendered.rstrip("0").rstrip(".")
            return rendered
        except InvalidOperation:
            pass
    return text


def normalize_tecdoc(value) -> str:
    """TECDOC arrives as text or as a float from Excel; normalize to digits."""
    if value is None:
        return ""
    text = _plain_number(value).strip()
    if not text:
        return ""
    if text.upper() in _MISSING_MARKERS:
        return ""
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text.upper()


def normalize_product(value) -> str:
    """Normalize SAP product identifiers exported as Excel numbers."""
    return normalize_tecdoc(value)


def normalize_code(value) -> str:
    """Tariff codes must reach SAP without separators or trailing float noise."""
    if value is None:
        return ""
    text = _plain_number(value).strip()
    if not text or text.upper() in _MISSING_MARKERS:
        return ""
    if text.endswith(".0") and text[:-2].replace(" ", "").isdigit():
        text = text[:-2]
    return "".join(ch for ch in text if ch.isalnum()).upper()


class HsReference:
    """TECDOC -> tariff code lookup built from the reference workbook."""

    def __init__(self, path: str | Path, sheet: str | None = None):
        self.path = str(path)
        self.sheet = sheet
        self.headers: list[str] = []
        self.header_row = 0
        self.rows_by_tecdoc: dict[str, dict] = {}
        self._header_lookup: dict[str, str] = {}
        self._values_by_tecdoc: dict[str, dict[int, object]] = {}
        self._value_conflicts: dict[str, dict[int, set[str]]] = {}
        self._load()

    def _load(self) -> None:
        wb = load_workbook(self.path, data_only=True, read_only=True)
        try:
            ws = wb[self.sheet] if self.sheet else wb[wb.sheetnames[0]]
            rows = ws.iter_rows(values_only=True)
            preview: list[tuple] = []
            for _ in range(25):
                row = next(rows, None)
                if row is None:
                    break
                preview.append(row)
            if not preview:
                raise ValueError("Fisierul de referinta este gol.")

            tecdoc_aliases = {
                "tecdoc",
                "generic article number (tecdoc)",
                "generic article number tecdoc",
            }
            header_pos = -1
            tec_idx = -1
            for pos, candidate in enumerate(preview):
                keys = [_header_key(value) for value in candidate]
                match = next(
                    (idx for idx, key in enumerate(keys) if key in tecdoc_aliases),
                    None,
                )
                if match is not None:
                    header_pos = pos
                    tec_idx = match
                    break
            if header_pos < 0:
                raise ValueError(
                    "Fisierul de referinta nu contine un antet TECDOC in primele "
                    "25 de randuri."
                ) from None

            header = preview[header_pos]
            self.header_row = header_pos + 1
            self.headers = [
                re.sub(r"\s+", " ", str(h).replace("\ufeff", " ")).strip()
                if h is not None
                else ""
                for h in header
            ]
            for original in self.headers:
                if not original:
                    continue
                key = _header_key(original)
                if key in self._header_lookup:
                    raise ValueError(
                        f"Antet duplicat in referinta: {original!r} "
                        f"(rand {self.header_row})."
                    )
                self._header_lookup[key] = original

            tariff_keys = {_header_key(meta["column"]) for meta in SCHEMES.values()}
            conflicts: list[str] = []
            data_rows = chain(preview[header_pos + 1 :], rows)
            for row in data_rows:
                if row is None or len(row) <= tec_idx:
                    continue
                tecdoc = normalize_tecdoc(row[tec_idx])
                if not tecdoc:
                    continue
                record = {
                    self.headers[i]: row[i]
                    for i in range(min(len(self.headers), len(row)))
                    if self.headers[i]
                }
                values = {
                    i: value
                    for i, value in enumerate(row)
                    if value is not None and str(value).strip()
                }
                existing_values = self._values_by_tecdoc.setdefault(tecdoc, {})
                for idx, value in values.items():
                    current = existing_values.get(idx)
                    if current is None or not str(current).strip():
                        existing_values[idx] = value
                        continue
                    old_value = _plain_number(current).strip()
                    new_value = _plain_number(value).strip()
                    if old_value and new_value and old_value != new_value:
                        conflicts_for_tecdoc = self._value_conflicts.setdefault(
                            tecdoc, {}
                        )
                        conflicts_for_tecdoc.setdefault(idx, set()).update(
                            (old_value, new_value)
                        )
                existing = self.rows_by_tecdoc.get(tecdoc)
                if existing is None:
                    self.rows_by_tecdoc[tecdoc] = record
                    continue

                # Duplicate TECDOC rows are common in hand-maintained files.
                # Merge complementary cells, but reject two different tariff
                # codes for the same scheme: choosing either silently is unsafe.
                for column, value in record.items():
                    current = existing.get(column)
                    current_blank = current is None or not str(current).strip()
                    value_blank = value is None or not str(value).strip()
                    if current_blank and not value_blank:
                        existing[column] = value
                        continue
                    if _header_key(column) not in tariff_keys:
                        continue
                    old_code = normalize_code(current)
                    new_code = normalize_code(value)
                    if old_code and new_code and old_code != new_code:
                        conflicts.append(
                            f"TECDOC {tecdoc}, coloana {column}: "
                            f"{old_code} vs {new_code}"
                        )
            if conflicts:
                raise ValueError(
                    "Referinta contine coduri tarifare conflictuale pentru acelasi "
                    "TECDOC: " + " | ".join(conflicts[:5])
                )
        finally:
            wb.close()

    def has_column(self, column: str) -> bool:
        return _header_key(column) in self._header_lookup

    def code_for(self, tecdoc: str, column: str) -> str:
        record = self.rows_by_tecdoc.get(normalize_tecdoc(tecdoc))
        if not record:
            return ""
        actual = self._header_lookup.get(_header_key(column))
        return normalize_code(record.get(actual)) if actual else ""

    def _column_index(self, column: str | int) -> int:
        """Resolve an Excel letter or a header name to a zero-based index."""
        if isinstance(column, int):
            if column < 1:
                raise ValueError(f"Coloana Excel invalida: {column!r}")
            return column - 1

        requested = str(column or "").strip()
        if not requested:
            raise ValueError("Coloana Excel pentru descriere nu este configurata.")
        if re.fullmatch(r"[A-Za-z]{1,3}", requested):
            return column_index_from_string(requested.upper()) - 1

        actual = self._header_lookup.get(_header_key(requested))
        if actual is None:
            raise ValueError(f"Referinta nu contine coloana/antetul {requested!r}.")
        return self.headers.index(actual)

    def value_for(self, tecdoc: str, column: str | int) -> str:
        """Read a reference cell by TECDOC and Excel letter/header.

        Conflicting duplicate TECDOC rows are rejected for the requested
        column. This is especially important for commercial descriptions: the
        automation must never choose one of two different texts silently.
        """
        normalized = normalize_tecdoc(tecdoc)
        if not normalized or normalized not in self.rows_by_tecdoc:
            return ""
        idx = self._column_index(column)
        conflicts = self._value_conflicts.get(normalized, {}).get(idx) or set()
        if conflicts:
            shown = " vs ".join(sorted(conflicts)[:3])
            raise ValueError(
                f"TECDOC {normalized} are valori conflictuale in coloana "
                f"{column}: {shown}"
            )
        value = self._values_by_tecdoc.get(normalized, {}).get(idx)
        if value is None:
            return ""
        text = str(value).strip()
        return "" if text.upper() in _MISSING_MARKERS else text

    def commercial_descriptions(
        self, tecdoc: str, languages: list[dict]
    ) -> dict[str, str]:
        """Return configured language descriptions for a TECDOC row."""
        descriptions: dict[str, str] = {}
        for language in languages:
            code = str(language.get("code") or "").strip().upper()
            column = language.get("column")
            if not code or not column:
                continue
            text = self.value_for(tecdoc, column)
            if text:
                descriptions[code] = text
        return descriptions

    def validate_description_columns(self, languages: list[dict]) -> None:
        """Fail early when a configured description column cannot exist."""
        for language in languages:
            code = str(language.get("code") or "").strip().upper() or "?"
            column = language.get("column")
            idx = self._column_index(column)
            if idx >= len(self.headers):
                raise ValueError(
                    f"Coloana {column!r} pentru descrierea {code} nu exista "
                    f"in referinta (ultima coloana: {len(self.headers)})."
                )

    def describe(self, tecdoc: str) -> str:
        record = self.rows_by_tecdoc.get(normalize_tecdoc(tecdoc))
        if not record:
            return ""
        for candidate in ("TECDOC text", "Text ONR", "Generic Article Text"):
            actual = self._header_lookup.get(_header_key(candidate))
            if actual and record.get(actual) not in (None, ""):
                return str(record[actual]).strip()
        return ""

    @property
    def tecdoc_count(self) -> int:
        return len(self.rows_by_tecdoc)
