"""Parse a SAP GTS 'Classify Products' worklist export and decide what can be
maintained.

A product is maintainable when its TECDOC number resolves to a tariff code in
the reference workbook for the selected numbering scheme. Products sharing a
TECDOC share the HS code, so they are grouped and can be mass-classified in one
SAP dialog.
"""

from __future__ import annotations

import re
from itertools import chain
from pathlib import Path

from hs_reference import (
    HsReference,
    normalize_code,
    normalize_product,
    normalize_tecdoc,
    scheme_info,
)
from openpyxl import load_workbook

# Header aliases, matched case-insensitively against the exported sheet.
_PRODUCT_HEADERS = ("product number", "material number", "product", "material")
_TECDOC_HEADERS = ("generic article number (tecdoc)", "tecdoc")
_TECDOC_TEXT_HEADERS = ("text generic article number (tecdoc)", "tecdoc text")
_TEXT_HEADERS = ("product short text", "short text", "description")
_MATERIAL_TYPE_HEADERS = ("material type",)
_COUNTRY_HEADERS = ("country/country group", "country")
_STATUS_HEADERS = ("product classification status", "cust.product sts")
_COUNTRY_LANGUAGE = {"EU": "DE", "GB": "EN"}

REASON_LABELS = {
    "missing_tecdoc": "Fara numar TECDOC in worklist",
    "conflicting_tecdoc": "Produsul are numere TECDOC diferite in worklist",
    "tecdoc_not_in_reference": "TECDOC lipseste din fisierul de referinta",
    "no_code_for_scheme": "Fara cod tarifar pentru schema selectata",
}


def _find_index(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    lowered = [
        re.sub(r"\s+", " ", h.replace("\ufeff", " ").replace("\xa0", " "))
        .strip()
        .casefold()
        for h in headers
    ]
    for candidate in candidates:
        if candidate in lowered:
            return lowered.index(candidate)
    # Prefer the most specific candidate. A header such as "Product
    # Classification Status" must never be mistaken for the Product number.
    for candidate in candidates:
        for idx, head in enumerate(lowered):
            if not head.startswith(candidate):
                continue
            if candidate in ("product", "material") and re.search(
                r"\b(status|classification|text|description|type)\b", head
            ):
                continue
            return idx
    return None


def country_groups_from_value(value) -> list[str]:
    """Return EU/GB tokens from the SAP Country/Country Group cell."""
    tokens = re.findall(r"[A-Za-z]+", str(value or "").upper())
    groups: list[str] = []
    for token in tokens:
        if token in _COUNTRY_LANGUAGE and token not in groups:
            groups.append(token)
    return groups


def description_languages_for_countries(countries: list[str] | str | None) -> list[str]:
    """Map worklist country groups to commercial-description languages.

    EU -> German (DE), GB -> English (EN). Anything else is ignored.
    """
    if isinstance(countries, str):
        groups = country_groups_from_value(countries)
    else:
        groups = []
        for value in countries or []:
            for token in country_groups_from_value(value):
                if token not in groups:
                    groups.append(token)
    languages: list[str] = []
    for group in groups:
        code = _COUNTRY_LANGUAGE[group]
        if code not in languages:
            languages.append(code)
    return languages


def read_worklist(path: str | Path, sheet: str | None = None) -> list[dict]:
    """Return raw worklist rows with the columns this tool needs."""
    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        preview: list[tuple] = []
        for _ in range(25):
            row = next(rows, None)
            if row is None:
                break
            preview.append(row)
        if not preview:
            raise ValueError("Exportul SAP este gol.")

        header_pos = -1
        headers: list[str] = []
        for pos, candidate in enumerate(preview):
            candidate_headers = [
                re.sub(r"\s+", " ", str(h).replace("\ufeff", " ")).strip()
                if h is not None
                else ""
                for h in candidate
            ]
            if (
                _find_index(candidate_headers, _PRODUCT_HEADERS) is not None
                and _find_index(candidate_headers, _TECDOC_HEADERS) is not None
            ):
                header_pos = pos
                headers = candidate_headers
                break
        if header_pos < 0:
            raise ValueError(
                "Exportul SAP nu contine antetele Product si Generic Article "
                "Number (TECDOC) in primele 25 de randuri."
            )

        idx_product = _find_index(headers, _PRODUCT_HEADERS)
        idx_tecdoc = _find_index(headers, _TECDOC_HEADERS)
        if idx_product is None:
            raise ValueError("Exportul SAP nu contine coloana 'Product'.")
        if idx_tecdoc is None:
            raise ValueError(
                "Exportul SAP nu contine coloana 'Generic Article Number (TECDOC)'."
            )
        idx_tecdoc_text = _find_index(headers, _TECDOC_TEXT_HEADERS)
        idx_text = _find_index(headers, _TEXT_HEADERS)
        idx_material_type = _find_index(headers, _MATERIAL_TYPE_HEADERS)
        idx_country = _find_index(headers, _COUNTRY_HEADERS)
        idx_status = _find_index(headers, _STATUS_HEADERS)

        def cell(row, idx):
            if idx is None or idx >= len(row):
                return ""
            value = row[idx]
            return "" if value is None else str(value).strip()

        items: list[dict] = []
        data_rows = chain(preview[header_pos + 1 :], rows)
        for excel_row, row in enumerate(data_rows, start=header_pos + 2):
            if row is None:
                continue
            product = normalize_product(cell(row, idx_product))
            if not product:
                continue
            items.append(
                {
                    "row": excel_row,
                    "product": product,
                    "text": cell(row, idx_text),
                    "material_type": cell(row, idx_material_type),
                    "tecdoc": normalize_tecdoc(cell(row, idx_tecdoc)),
                    "tecdoc_text": cell(row, idx_tecdoc_text),
                    "country": cell(row, idx_country),
                    "status": cell(row, idx_status),
                }
            )
        return items
    finally:
        wb.close()


def save_remaining_worklist(
    source_path: str | Path,
    target_path: str | Path,
    maintained_products: list[str] | set[str] | tuple[str, ...],
    sheet: str | None = None,
) -> dict:
    """Copy a SAP worklist and remove only SAP-confirmed product rows.

    The workbook is opened without ``data_only`` so formulas, formatting,
    column order, preamble rows, and any additional sheets are retained. Every
    occurrence of a confirmed Product is removed from the worklist sheet; all
    other rows remain in their original order.
    """
    source = Path(source_path)
    target = Path(target_path)
    keep_vba = source.suffix.lower() == ".xlsm"
    wb = load_workbook(source, keep_vba=keep_vba)
    try:
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        preview = list(
            ws.iter_rows(
                min_row=1,
                max_row=min(ws.max_row, 25),
                values_only=True,
            )
        )
        header_pos = -1
        headers: list[str] = []
        for pos, candidate in enumerate(preview):
            candidate_headers = [
                re.sub(r"\s+", " ", str(value).replace("\ufeff", " ")).strip()
                if value is not None
                else ""
                for value in candidate
            ]
            if _find_index(candidate_headers, _PRODUCT_HEADERS) is not None:
                header_pos = pos
                headers = candidate_headers
                break
        if header_pos < 0:
            raise ValueError(
                "Exportul SAP nu contine antetul Product in primele 25 de randuri."
            )
        product_index = _find_index(headers, _PRODUCT_HEADERS)
        if product_index is None:
            raise ValueError("Exportul SAP nu contine coloana 'Product'.")

        maintained = {
            normalize_product(product)
            for product in maintained_products
            if normalize_product(product)
        }
        first_data_row = header_pos + 2
        product_column = product_index + 1
        original_products: set[str] = set()
        removed_products: set[str] = set()
        rows_to_delete: list[int] = []
        for row_number in range(first_data_row, ws.max_row + 1):
            product = normalize_product(ws.cell(row_number, product_column).value)
            if not product:
                continue
            original_products.add(product)
            if product in maintained:
                rows_to_delete.append(row_number)
                removed_products.add(product)

        # Delete contiguous ranges from bottom to top. This is materially faster
        # than deleting thousands of individual rows and does not disturb the
        # row numbers still waiting to be removed.
        ranges: list[tuple[int, int]] = []
        for row_number in rows_to_delete:
            if ranges and row_number == ranges[-1][1] + 1:
                ranges[-1] = (ranges[-1][0], row_number)
            else:
                ranges.append((row_number, row_number))
        for start, end in reversed(ranges):
            ws.delete_rows(start, end - start + 1)

        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
        remaining_products = original_products - removed_products
        unmatched_products = maintained - removed_products
        return {
            "source": source.name,
            "file": target.name,
            "sheet": ws.title,
            "header_row": header_pos + 1,
            "original_products": len(original_products),
            "requested_products": len(maintained),
            "removed_products": len(removed_products),
            "removed_rows": len(rows_to_delete),
            "remaining_products": len(remaining_products),
            "unmatched_products": sorted(unmatched_products),
        }
    finally:
        wb.close()


def analyze(
    worklist: list[dict],
    reference: HsReference,
    scheme: str,
) -> dict:
    """Split the worklist into maintainable groups and blocked products."""
    info = scheme_info(scheme)
    column = info["column"]
    if not reference.has_column(column):
        raise ValueError(
            f"Fisierul de referinta nu contine coloana '{column}' "
            f"ceruta de schema {info['code']}."
        )

    ready: list[dict] = []
    blocked: list[dict] = []
    rows_by_product: dict[str, list[dict]] = {}
    for item in worklist:
        rows_by_product.setdefault(item["product"], []).append(item)

    for product, candidates in rows_by_product.items():
        tecdocs = list(
            dict.fromkeys(item["tecdoc"] for item in candidates if item.get("tecdoc"))
        )
        # Prefer the populated duplicate when another row for the same product
        # has an empty TecDoc. Different non-empty values are ambiguous and must
        # be reviewed instead of silently taking the first export row.
        item = next((row for row in candidates if row.get("tecdoc")), candidates[0])
        record = dict(item)
        record["scheme"] = info["code"]
        country_groups: list[str] = []
        for row in candidates:
            for token in country_groups_from_value(row.get("country") or ""):
                if token not in country_groups:
                    country_groups.append(token)
        record["country_groups"] = country_groups
        record["country"] = " ".join(country_groups) or str(item.get("country") or "")

        if len(tecdocs) > 1:
            record["tecdoc"] = " | ".join(tecdocs)
            record["reason"] = "conflicting_tecdoc"
            blocked.append(record)
            continue
        if not item["tecdoc"]:
            record["reason"] = "missing_tecdoc"
            blocked.append(record)
            continue
        if item["tecdoc"] not in reference.rows_by_tecdoc:
            record["reason"] = "tecdoc_not_in_reference"
            blocked.append(record)
            continue

        code = reference.code_for(item["tecdoc"], column)
        if not code:
            record["reason"] = "no_code_for_scheme"
            blocked.append(record)
            continue

        record["hs_code"] = code
        record["reference_text"] = reference.describe(item["tecdoc"])
        ready.append(record)

    groups: dict[str, dict] = {}
    for record in ready:
        group = groups.setdefault(
            record["hs_code"],
            {
                "hs_code": record["hs_code"],
                "scheme": info["code"],
                "tecdocs": [],
                "products": [],
            },
        )
        if record["tecdoc"] not in group["tecdocs"]:
            group["tecdocs"].append(record["tecdoc"])
        group["products"].append(
            {
                "product": record["product"],
                "text": record["text"],
                "material_type": record.get("material_type", ""),
                "tecdoc": record["tecdoc"],
                "tecdoc_text": record["tecdoc_text"],
                "country": record.get("country", ""),
                "country_groups": list(record.get("country_groups") or []),
            }
        )

    ordered_groups = sorted(
        groups.values(), key=lambda g: (-len(g["products"]), g["hs_code"])
    )
    for group in ordered_groups:
        group["count"] = len(group["products"])

    blocked_by_reason: dict[str, int] = {}
    for record in blocked:
        blocked_by_reason[record["reason"]] = (
            blocked_by_reason.get(record["reason"], 0) + 1
        )

    return {
        "scheme": info["code"],
        "scheme_label": info["label"],
        "reference_column": column,
        "variant": info["variant"],
        "summary": {
            "worklist_rows": len(worklist),
            "distinct_products": len(rows_by_product),
            "ready": len(ready),
            "blocked": len(blocked),
            "groups": len(ordered_groups),
            "blocked_by_reason": blocked_by_reason,
        },
        "groups": ordered_groups,
        "ready": ready,
        "blocked": blocked,
    }


def blocked_workbook_rows(blocked: list[dict]) -> list[list[str]]:
    """Rows for the 'not ready' export the operator reviews separately."""
    rows = [
        [
            "product",
            "product_text",
            "material_type",
            "tecdoc",
            "tecdoc_text",
            "country",
            "scheme",
            "reason",
        ]
    ]
    for item in blocked:
        rows.append(
            [
                item.get("product", ""),
                item.get("text", ""),
                item.get("material_type", ""),
                item.get("tecdoc", ""),
                item.get("tecdoc_text", ""),
                item.get("country", ""),
                item.get("scheme", ""),
                REASON_LABELS.get(item.get("reason", ""), item.get("reason", "")),
            ]
        )
    return rows


def selected_groups(analysis: dict, hs_codes: list[str] | None) -> list[dict]:
    """Restrict an analysis to the HS codes the operator approved."""
    if hs_codes is None:
        return list(analysis.get("groups") or [])
    wanted = {normalize_code(code) for code in hs_codes if code}
    if not wanted:
        return []
    return [g for g in analysis.get("groups") or [] if g["hs_code"] in wanted]


def build_description_plan(
    groups: list[dict],
    reference: HsReference,
    languages: list[dict],
) -> dict:
    """Build commercial-description work from the SAP worklist country group.

    EU keeps German (reference DE), GB keeps English (reference EN). A product
    without EU/GB is skipped. Only the languages required by those country
    groups must exist in the TECDOC reference before any HS write starts.
    """
    language_by_code = {
        str(language.get("code") or "").strip().upper(): language
        for language in languages
        if language.get("code") and language.get("column")
    }
    if not language_by_code:
        raise ValueError("Nu sunt configurate limbile pentru descrieri comerciale.")

    reference.validate_description_columns(languages)
    items: list[dict] = []
    missing: list[dict] = []
    skipped: list[dict] = []
    seen_products: dict[str, str] = {}
    planned_languages: list[str] = []

    for group in groups:
        for row in group.get("products") or []:
            product = normalize_product(row.get("product"))
            tecdoc = normalize_tecdoc(row.get("tecdoc"))
            if not product:
                continue
            previous_tecdoc = seen_products.get(product)
            if previous_tecdoc and previous_tecdoc != tecdoc:
                raise ValueError(
                    f"Produsul {product} are TECDOC-uri diferite in planul "
                    f"descrierilor: {previous_tecdoc} vs {tecdoc}."
                )
            if previous_tecdoc:
                continue
            seen_products[product] = tecdoc

            country_groups = list(row.get("country_groups") or [])
            if not country_groups:
                country_groups = country_groups_from_value(row.get("country") or "")
            needed_codes = description_languages_for_countries(country_groups)
            if not needed_codes:
                skipped.append(
                    {
                        "material": product,
                        "product": product,
                        "tecdoc": tecdoc,
                        "country_groups": country_groups,
                    }
                )
                continue

            configured = [
                language_by_code[code]
                for code in needed_codes
                if code in language_by_code
            ]
            if len(configured) != len(needed_codes):
                missing.append(
                    {
                        "material": product,
                        "product": product,
                        "tecdoc": tecdoc,
                        "country_groups": country_groups,
                        "missing_languages": [
                            code
                            for code in needed_codes
                            if code not in language_by_code
                        ],
                    }
                )
                continue

            descriptions = reference.commercial_descriptions(tecdoc, configured)
            missing_codes = [
                code for code in needed_codes if not descriptions.get(code)
            ]
            if not tecdoc or missing_codes:
                missing.append(
                    {
                        "material": product,
                        "product": product,
                        "tecdoc": tecdoc,
                        "country_groups": country_groups,
                        "missing_languages": missing_codes or needed_codes,
                    }
                )
                continue
            for code in needed_codes:
                if code not in planned_languages:
                    planned_languages.append(code)
            items.append(
                {
                    "material": product,
                    "product": product,
                    "tecdoc": tecdoc,
                    "country_groups": country_groups,
                    "languages": needed_codes,
                    "descriptions": {code: descriptions[code] for code in needed_codes},
                }
            )

    return {
        "languages": planned_languages,
        "items": items,
        "missing": missing,
        "skipped": skipped,
        "summary": {
            "products": len(seen_products),
            "ready": len(items),
            "missing": len(missing),
            "skipped": len(skipped),
        },
    }
