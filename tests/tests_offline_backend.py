"""Offline regression for the HS backend: config, analysis, and route guards.

Uses the Flask test client. No SAP session and no browser are started.
"""

import json
import os
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "server"))

import app as srv


class FakeSession:
    """Stands in for the SAP session so no browser is launched."""

    def __init__(self, state: str):
        self.state = {"state": state, "msg": ""}
        self.started: dict | None = None
        self.submitted: list[str] = []
        self.payload: dict = {}

    def start(self, cfg, creds, cb):
        self.started = {"cfg": cfg, "creds": creds}
        self.state = {"state": "starting", "msg": ""}

    def submit(self, command, payload, cb, cancel):
        self.submitted.append(command)
        self.payload = payload

    def stop(self):
        self.state = {"state": "stopped", "msg": "Browser inchis."}


def case_shutdown_is_local_only(client) -> bool:
    """Shutdown must stay local and must not exit the test process."""
    original_exit = srv._schedule_process_exit
    original_flag = srv._SHUTDOWN_SCHEDULED
    original_hs_session = srv.get_hs_session
    original_session = srv.get_session
    original_hs_cleanup = srv.cleanup_hs_lock
    original_desc_cleanup = srv.cleanup_user_data_lock
    scheduled: list[float] = []
    fake = FakeSession("ready")

    def fake_exit(delay_s: float = 0.4) -> None:
        scheduled.append(delay_s)

    srv._schedule_process_exit = fake_exit
    srv._SHUTDOWN_SCHEDULED = False
    srv.get_hs_session = lambda: fake
    srv.get_session = lambda: fake
    srv.cleanup_hs_lock = lambda kill_chrome=True: {
        "killed_pids": [11] if kill_chrome else [],
        "removed": [],
    }
    srv.cleanup_user_data_lock = lambda kill_chrome=True: {
        "killed_pids": [22] if kill_chrome else [],
        "removed": [],
    }
    try:
        remote = client.post("/shutdown", environ_base={"REMOTE_ADDR": "8.8.8.8"})
        local = client.post("/shutdown")
        again = client.post("/shutdown")
    finally:
        srv._schedule_process_exit = original_exit
        srv._SHUTDOWN_SCHEDULED = original_flag
        srv.get_hs_session = original_hs_session
        srv.get_session = original_session
        srv.cleanup_hs_lock = original_hs_cleanup
        srv.cleanup_user_data_lock = original_desc_cleanup

    remote_blocked = remote.status_code == 403
    local_ok = (
        local.status_code == 200
        and local.get_json().get("shutting_down") is True
        and len(scheduled) == 1
    )
    idempotent = again.status_code == 200 and len(scheduled) == 1
    ok = remote_blocked and local_ok and idempotent
    print(
        f"SHUTDOWN_LOCAL: remote={remote.status_code} local={local.status_code} "
        f"scheduled={len(scheduled)} -> {'PASS' if ok else 'FAIL'}"
    )
    return ok


def case_startup_failure_is_terminal(client) -> bool:
    """A worker startup failure must close SSE and expose a report."""
    original_session = srv.get_session

    class FailingSession:
        state = {"state": "stopped", "msg": ""}

        def start(self, cfg, creds, cb):
            self.state = {"state": "starting", "msg": "Pornesc browserul..."}
            return {"ok": True, "state": "starting"}

        def run_items(self, items, cfg, event_cb, cancel_event):
            event_cb({"type": "fatal", "error": "startup failure"})
            event_cb({"type": "_end_"})

    srv.get_session = lambda: FailingSession()
    try:
        response = client.post(
            "/run",
            data=json.dumps(
                {
                    "items": [{"material": "100001", "description": "Text"}],
                    "credentials": None,
                }
            ),
            content_type="application/json",
        )
        job = response.get_json()["job"]
        events = client.get(f"/events/{job}").get_data(as_text=True)
        report = client.get(f"/report/{job}")
        data_rows = report.get_data(as_text=True)
    finally:
        srv.get_session = original_session

    terminal = 'data: {"type": "fatal", "error": "startup failure"}' in events
    ended = events.count("event: end") == 1
    report_ready = report.status_code == 200 and "FAIL" in data_rows
    ok = response.status_code == 200 and terminal and ended and report_ready
    print(
        f"STARTUP_FAILURE_TERMINAL: response={response.status_code} "
        f"terminal={terminal} ended={ended} report={report.status_code} "
        f"-> {'PASS' if ok else 'FAIL'}"
    )
    return ok


def case_reference_autoload_reload() -> bool:
    """The startup cache must refresh after an operator updates the workbook."""
    with tempfile.TemporaryDirectory(prefix="hs-reference-cache-") as tmp:
        path = Path(tmp) / "reference.xlsx"

        def save(rows):
            wb = Workbook()
            ws = wb.active
            for row in rows:
                ws.append(row)
            wb.save(path)
            wb.close()

        header = [
            "TECDOC",
            "Text",
            "EU TARIC",
            None,
            None,
            None,
            None,
            None,
            None,
            "German",
            "English",
        ]
        save(
            [
                header,
                [1, "One", "1111", None, None, None, None, None, None, "DE 1", "EN 1"],
            ]
        )
        cfg = {
            "reference_file": str(path),
            "description_languages": [
                {"code": "DE", "column": "J", "country_name": "Germany"},
                {
                    "code": "EN",
                    "column": "K",
                    "country_name": "United Kingdom",
                },
            ],
        }
        first = srv.load_reference(cfg)
        second = srv.load_reference(cfg)
        cached = first is second and first.tecdoc_count == 1

        previous_mtime = path.stat().st_mtime_ns
        save(
            [
                header,
                [1, "One", "1111", None, None, None, None, None, None, "DE 1", "EN 1"],
                [2, "Two", "2222", None, None, None, None, None, None, "DE 2", "EN 2"],
            ]
        )
        current = path.stat()
        if current.st_mtime_ns <= previous_mtime:
            os.utime(
                path,
                ns=(current.st_atime_ns, previous_mtime + 1_000_000),
            )
        refreshed = srv.load_reference(cfg)
        status = srv.reference_status(cfg)
        reloaded = (
            refreshed is not first
            and refreshed.tecdoc_count == 2
            and status["reference_loaded"]
            and status["reference_tecdoc_count"] == 2
        )
        ok = cached and reloaded
        print(
            f"REFERENCE_AUTOLOAD: cached={cached} reloaded={reloaded} -> "
            f"{'PASS' if ok else 'FAIL'}"
        )
        return ok


def case_description_preflight_guard(client) -> bool:
    """Missing DE/EN texts are recorded; classification still starts."""
    with tempfile.TemporaryDirectory(prefix="hs-description-guard-") as tmp:
        tmp_path = Path(tmp)

        def save(path: Path, english: str | None):
            wb = Workbook()
            ws = wb.active
            ws.append(
                [
                    "TECDOC",
                    "Text",
                    "EU TARIC",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "German",
                    "English",
                ]
            )
            ws.append(
                [
                    1001,
                    "Sensor",
                    "90328900",
                    None,
                    None,
                    None,
                    None,
                    None,
                    None,
                    "Sensor DE",
                    english,
                ]
            )
            wb.save(path)
            wb.close()

        missing_path = tmp_path / "missing.xlsx"
        complete_path = tmp_path / "complete.xlsx"
        save(missing_path, None)
        save(complete_path, "Sensor EN")

        result = {
            "scheme": "EDCHSCDEEX",
            "groups": [
                {
                    "hs_code": "90328900",
                    "count": 1,
                    "products": [
                        {
                            "product": "355135851",
                            "tecdoc": "1001",
                            "country": "GB",
                        }
                    ],
                }
            ],
            "blocked": [],
        }
        analysis_id = srv.store_analysis(result, "synthetic.xlsx")
        languages = [
            {"code": "DE", "column": "J", "country_name": "Germany"},
            {
                "code": "EN",
                "column": "K",
                "country_name": "United Kingdom",
            },
        ]
        active_path = missing_path

        def fake_hs_config(*args, **kwargs):
            return {
                "allow_commit": True,
                "reference_file": str(active_path),
                "reference_sheet": None,
                "description_languages": languages,
                "description_display_maintained": False,
                "sap_url": "https://cift6.example/sap/bc/ui2/flp?client=100"
                "#CustomsProduct-classify?sap-ui-tech-hint=GUI",
            }

        original_hs_config = srv.hs_config
        original_session = srv.get_hs_session
        fake = FakeSession("ready")
        try:
            srv.hs_config = fake_hs_config
            srv.get_hs_session = lambda: fake
            missing = client.post(
                "/hs/classify",
                data=json.dumps(
                    {
                        "analysis": analysis_id,
                        "hs_codes": ["90328900"],
                        "allow_commit": True,
                    }
                ),
                content_type="application/json",
            )
            missing_plan = fake.payload.get("description_plan") or {}
            missing_queued = (
                missing.status_code == 200
                and fake.submitted == ["classify"]
                and (missing_plan.get("summary") or {}).get("missing") == 1
            )
            fake.submitted = []
            fake.payload = {}

            active_path = complete_path
            complete = client.post(
                "/hs/classify",
                data=json.dumps(
                    {
                        "analysis": analysis_id,
                        "hs_codes": ["90328900"],
                        "allow_commit": True,
                    }
                ),
                content_type="application/json",
            )
            plan = fake.payload.get("description_plan") or {}
            complete_queued = (
                complete.status_code == 200
                and fake.submitted == ["classify"]
                and fake.payload.get("description_display_maintained") is False
                and plan.get("summary")
                == {"products": 1, "ready": 1, "missing": 0, "skipped": 0}
                and plan.get("items", [])[0].get("languages") == ["EN"]
                and plan.get("items", [])[0].get("descriptions") == {"EN": "Sensor EN"}
            )
        finally:
            srv.hs_config = original_hs_config
            srv.get_hs_session = original_session

        ok = missing_queued and complete_queued
        print(
            f"DESCRIPTION_PREFLIGHT: missing_queued={missing_queued} "
            f"complete_queued={complete_queued} -> {'PASS' if ok else 'FAIL'}"
        )
        return ok


def case_manual_descriptions_route(client) -> bool:
    """Selected worklist groups can start DE/EN without rewriting HS."""
    with tempfile.TemporaryDirectory(prefix="hs-desc-only-") as tmp:
        path = Path(tmp) / "complete.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.append(
            [
                "TECDOC",
                "Text",
                "EU TARIC",
                None,
                None,
                None,
                None,
                None,
                None,
                "German",
                "English",
            ]
        )
        ws.append(
            [
                1001,
                "Sensor",
                "90328900",
                None,
                None,
                None,
                None,
                None,
                None,
                "Sensor DE",
                "Sensor EN",
            ]
        )
        wb.save(path)
        wb.close()

        result = {
            "scheme": "EDCHSCDEEX",
            "groups": [
                {
                    "hs_code": "90328900",
                    "count": 1,
                    "products": [
                        {
                            "product": "355135851",
                            "tecdoc": "1001",
                            "country": "GB",
                        }
                    ],
                }
            ],
            "blocked": [],
        }
        analysis_id = srv.store_analysis(result, "synthetic.xlsx")
        languages = [
            {"code": "DE", "column": "J", "country_name": "Germany"},
            {"code": "EN", "column": "K", "country_name": "United Kingdom"},
        ]

        def fake_hs_config(*args, **kwargs):
            return {
                "allow_commit": True,
                "reference_file": str(path),
                "reference_sheet": None,
                "description_languages": languages,
                "description_display_maintained": False,
                "sap_url": "https://cift6.example/sap/bc/ui2/flp?client=100"
                "#CustomsProduct-classify?sap-ui-tech-hint=GUI",
            }

        original_hs_config = srv.hs_config
        original_session = srv.get_hs_session
        fake = FakeSession("ready")
        try:
            srv.hs_config = fake_hs_config
            srv.get_hs_session = lambda: fake
            dry = client.post(
                "/hs/descriptions",
                data=json.dumps(
                    {
                        "analysis": analysis_id,
                        "hs_codes": ["90328900"],
                    }
                ),
                content_type="application/json",
            )
            started = client.post(
                "/hs/descriptions",
                data=json.dumps(
                    {
                        "analysis": analysis_id,
                        "hs_codes": ["90328900"],
                        "allow_commit": True,
                    }
                ),
                content_type="application/json",
            )
            body = started.get_json() or {}
        finally:
            srv.hs_config = original_hs_config
            srv.get_hs_session = original_session

        ok = (
            dry.status_code == 400
            and started.status_code == 200
            and fake.submitted == ["descriptions"]
            and body.get("descriptions_only") is True
            and body.get("description_products") == 1
            and "groups" not in fake.payload
        )
        print(
            f"MANUAL_DESCRIPTIONS: dry={dry.status_code} start={started.status_code} "
            f"cmd={fake.submitted} -> {'PASS' if ok else 'FAIL'}"
        )
        return ok


def case_commit_guard(client) -> bool:
    """Commit requires both the request flag and the server-side gate."""
    result = {
        "scheme": "EDCHSCDEEX",
        "groups": [
            {
                "hs_code": "90328900",
                "count": 1,
                "products": [{"product": "355135851"}],
            }
        ],
        "blocked": [],
    }
    analysis_id = srv.store_analysis(result, "synthetic.xlsx")
    original_session = srv.get_hs_session
    original_hs_config = srv.hs_config
    fake = FakeSession("ready")
    try:
        srv.get_hs_session = lambda: fake
        missing_selection = client.post(
            "/hs/classify",
            data=json.dumps({"analysis": analysis_id}),
            content_type="application/json",
        )
        empty_selection = client.post(
            "/hs/classify",
            data=json.dumps({"analysis": analysis_id, "hs_codes": []}),
            content_type="application/json",
        )
        selection_guard = (
            missing_selection.status_code == 400 and empty_selection.status_code == 400
        )

        # Missing per-run approval must remain a dry run even when the local
        # installation is configured to permit commit.
        response = client.post(
            "/hs/classify",
            data=json.dumps({"analysis": analysis_id, "hs_codes": ["90328900"]}),
            content_type="application/json",
        )
        request_guard = bool(response.get_json().get("dry_run")) and not bool(
            fake.payload.get("allow_commit")
        )

        fake.payload = {}
        response = client.post(
            "/hs/classify",
            data=json.dumps(
                {
                    "analysis": analysis_id,
                    "hs_codes": ["90328900"],
                    "allow_commit": "true",
                }
            ),
            content_type="application/json",
        )
        type_guard = bool(response.get_json().get("dry_run")) and not bool(
            fake.payload.get("allow_commit")
        )

        # A client cannot override a disabled server-side gate.
        fake.payload = {}

        def commit_disabled(*args, **kwargs):
            cfg = original_hs_config(*args, **kwargs)
            cfg["allow_commit"] = False
            return cfg

        srv.hs_config = commit_disabled
        response = client.post(
            "/hs/classify",
            data=json.dumps(
                {
                    "analysis": analysis_id,
                    "hs_codes": ["90328900"],
                    "allow_commit": True,
                }
            ),
            content_type="application/json",
        )
        server_guard = bool(response.get_json().get("dry_run")) and not bool(
            fake.payload.get("allow_commit")
        )
    finally:
        srv.get_hs_session = original_session
        srv.hs_config = original_hs_config

    ok = (
        response.status_code == 200
        and selection_guard
        and request_guard
        and type_guard
        and server_guard
    )
    print(
        f"COMMIT_GUARD: selection={selection_guard} request={request_guard} "
        f"type={type_guard} "
        f"server={server_guard} -> "
        f"{'PASS' if ok else 'FAIL'}"
    )
    return ok


def case_semi_auto_starts_session(client) -> bool:
    """Semi-automatic maintenance must open SAP itself, not refuse the job."""
    sample = sorted((ROOT / "downloads").glob("worklist_*.xlsx"))
    if not sample:
        print("SEMI_AUTO_STARTS_SESSION: SKIP (niciun export in downloads/)")
        return True

    from hs_worklist import analyze, read_worklist

    hs = srv.hs_config()
    result = analyze(
        read_worklist(sample[-1]),
        srv.load_reference(hs),
        hs.get("default_scheme", "EDCHSCDEEX"),
    )
    if not result["groups"]:
        print("SEMI_AUTO_STARTS_SESSION: SKIP (niciun grup)")
        return True

    analysis_id = srv.store_analysis(result, sample[-1].name)
    fake = FakeSession("stopped")
    original = srv.get_hs_session
    srv.get_hs_session = lambda: fake
    try:
        r = client.post(
            "/hs/classify",
            data=json.dumps(
                {
                    "analysis": analysis_id,
                    "hs_codes": [result["groups"][0]["hs_code"]],
                    "headless": False,
                    "credentials": {"username": "u", "password": "p"},
                }
            ),
            content_type="application/json",
        )
        # No variant supplied: the scheme default must be filled in.
        fallback_variant = fake.payload.get("variant")

        explicit = FakeSession("stopped")
        srv.get_hs_session = lambda: explicit
        client.post(
            "/hs/classify",
            data=json.dumps(
                {
                    "analysis": analysis_id,
                    "hs_codes": [result["groups"][0]["hs_code"]],
                    "variant": "MY_VARIANT",
                }
            ),
            content_type="application/json",
        )
        explicit_variant = explicit.payload.get("variant")
    finally:
        srv.get_hs_session = original

    data = r.get_json()
    started = fake.started is not None
    visible = started and fake.started["cfg"].get("headless") is False
    queued = fake.submitted == ["classify"]
    dry_run = bool(data.get("dry_run"))
    variant_ok = bool(fallback_variant) and explicit_variant == "MY_VARIANT"
    ok = (
        r.status_code == 200
        and started
        and visible
        and queued
        and dry_run
        and variant_ok
    )
    print(
        f"SEMI_AUTO_STARTS_SESSION: {r.status_code} started={started} "
        f"visible={visible} queued={fake.submitted} dry_run={dry_run} -> "
        f"{'PASS' if ok else 'FAIL'}"
    )
    print(
        f"VARIANT_PASSED: default={fallback_variant!r} explicit={explicit_variant!r} "
        f"-> {'PASS' if variant_ok else 'FAIL'}"
    )
    return ok


def main() -> int:
    results: list[bool] = []
    report_tmp = tempfile.TemporaryDirectory(prefix="hs-backend-reports-")
    original_report_store = srv.HS_REPORT_STORE
    srv.HS_REPORT_STORE = srv.HsReportStore(Path(report_tmp.name))
    client = srv.app.test_client()

    def hs_cfg_now() -> dict:
        return srv.hs_config()

    # Commit uses a two-key guard: server configuration plus per-run approval.
    hs = srv.hs_config()
    results.append(case_startup_failure_is_terminal(client))
    results.append(case_shutdown_is_local_only(client))
    results.append(case_commit_guard(client))
    results.append(case_reference_autoload_reload())
    results.append(case_description_preflight_guard(client))
    results.append(case_manual_descriptions_route(client))

    r = client.post(
        "/hs/config",
        data=json.dumps({"allow_commit": "false"}),
        content_type="application/json",
    )
    strict_bool = r.status_code == 400
    print(f"COMMIT_CONFIG_TYPE: {r.status_code} -> {'PASS' if strict_bool else 'FAIL'}")
    results.append(strict_bool)

    original_display_maintained = hs.get("description_display_maintained", True)
    invalid_display = client.post(
        "/hs/config",
        data=json.dumps({"description_display_maintained": "false"}),
        content_type="application/json",
    )
    set_display = client.post(
        "/hs/config",
        data=json.dumps({"description_display_maintained": False}),
        content_type="application/json",
    )
    display_value = srv.hs_config().get("description_display_maintained")
    display_setting_ok = (
        invalid_display.status_code == 400
        and set_display.status_code == 200
        and display_value is False
    )
    client.post(
        "/hs/config",
        data=json.dumps(
            {"description_display_maintained": original_display_maintained}
        ),
        content_type="application/json",
    )
    print(
        "DESCRIPTION_DISPLAY_SETTING: "
        f"invalid={invalid_display.status_code} set={set_display.status_code} "
        f"-> {'PASS' if display_setting_ok else 'FAIL'}"
    )
    results.append(display_setting_ok)

    incomplete_languages = client.post(
        "/hs/config",
        data=json.dumps(
            {
                "description_languages": [
                    {"code": "DE", "column": "J", "country_name": "Germany"}
                ]
            }
        ),
        content_type="application/json",
    )
    language_guard = incomplete_languages.status_code == 400
    print(
        f"DESCRIPTION_LANGUAGE_GUARD: {incomplete_languages.status_code} -> "
        f"{'PASS' if language_guard else 'FAIL'}"
    )
    results.append(language_guard)

    r = client.get("/hs/schemes")
    schemes = r.get_json()
    ok = r.status_code == 200 and len(schemes) > 0
    print(
        f"SCHEMES: {r.status_code} n={len(schemes) if ok else '?'} -> "
        f"{'PASS' if ok else 'FAIL'}"
    )
    results.append(ok)

    # Launcher and both app pages must render.
    for path in ("/", "/apps/descriptions", "/apps/hs", "/apps/hs/reports"):
        r = client.get(path)
        ok = r.status_code == 200
        print(f"PAGE {path}: {r.status_code} -> {'PASS' if ok else 'FAIL'}")
        results.append(ok)

    # Classify must refuse an unknown analysis instead of touching SAP.
    r = client.post(
        "/hs/classify",
        data=json.dumps({"analysis": "nope", "hs_codes": ["1234"]}),
        content_type="application/json",
    )
    ok = r.status_code == 404
    print(f"CLASSIFY_UNKNOWN_ANALYSIS: {r.status_code} -> {'PASS' if ok else 'FAIL'}")
    results.append(ok)

    blocked_items = [
        {"product": f"P{i}", "reason": "missing_tecdoc", "text": ""} for i in range(80)
    ]
    analysis_id = srv.store_analysis(
        {
            "scheme": "EDCHSCDEEX",
            "groups": [
                {
                    "hs_code": "90328900",
                    "count": 1,
                    "products": [{"product": "355135851"}],
                }
            ],
            "blocked": blocked_items,
        },
        "synthetic.xlsx",
    )
    original_session = srv.get_hs_session
    fake = FakeSession("ready")
    try:
        srv.get_hs_session = lambda: fake
        response = client.post(
            "/hs/classify",
            data=json.dumps({"analysis": analysis_id, "hs_codes": ["90328900"]}),
            content_type="application/json",
        )
        payload_ok = response.status_code == 200 and fake.payload.get("blocked") == []
    finally:
        srv.get_hs_session = original_session
    print(
        f"CLASSIFY_SKIPS_BLOCKED_PAYLOAD: {response.status_code} "
        f"blocked={len(fake.payload.get('blocked') or [])} -> "
        f"{'PASS' if payload_ok else 'FAIL'}"
    )
    results.append(payload_ok)

    import hs_automation

    compact_summary = {
        "materials_blocked": 0,
        "blocked_products": [],
        "blocked_by_reason": {},
    }
    compact_logs: list[str] = []
    hs_automation._account_blocked_materials(
        compact_summary,
        blocked_items,
        "EDCHSCDEEX",
        lambda _level, msg: compact_logs.append(msg),
    )
    compact_ok = (
        compact_summary["materials_blocked"] == 80
        and len(compact_summary["blocked_products"]) == 50
        and len(compact_logs) == 1
    )
    print(
        f"BLOCKED_ACCOUNTING_COMPACT: blocked={compact_summary['materials_blocked']} "
        f"preview={len(compact_summary['blocked_products'])} logs={len(compact_logs)} "
        f"-> {'PASS' if compact_ok else 'FAIL'}"
    )
    results.append(compact_ok)

    # System switching is restricted to the configured allowlist.
    original = hs_cfg_now()["sap_url"]
    original_system = srv.hs_active_system(hs_cfg_now())

    r = client.post(
        "/hs/config",
        data=json.dumps({"system": "FT6"}),
        content_type="application/json",
    )
    data = r.get_json()
    ok = r.status_code == 200 and "cift6" in data.get("sap_url", "")
    print(
        f"SWITCH_FT6: {r.status_code} url={data.get('sap_url', '')[:40]} -> "
        f"{'PASS' if ok else 'FAIL'}"
    )
    results.append(ok)

    r = client.post(
        "/hs/config",
        data=json.dumps({"system": "FT1"}),
        content_type="application/json",
    )
    data = r.get_json()
    ok = r.status_code == 200 and "cift1" in data.get("sap_url", "")
    print(f"SWITCH_FT1: {r.status_code} -> {'PASS' if ok else 'FAIL'}")
    results.append(ok)

    r = client.post(
        "/hs/config",
        data=json.dumps({"system": "https://evil.example.com"}),
        content_type="application/json",
    )
    ok = r.status_code == 400
    print(f"SWITCH_REJECTS_UNKNOWN: {r.status_code} -> {'PASS' if ok else 'FAIL'}")
    results.append(ok)

    # A raw sap_url must not be settable any more.
    r = client.post(
        "/hs/config",
        data=json.dumps({"sap_url": "https://evil.example.com/x"}),
        content_type="application/json",
    )
    after = hs_cfg_now()["sap_url"]
    ok = "evil.example.com" not in after
    print(f"RAW_URL_IGNORED: url={after[:40]} -> {'PASS' if ok else 'FAIL'}")
    results.append(ok)

    if original_system:
        client.post(
            "/hs/config",
            data=json.dumps({"system": original_system}),
            content_type="application/json",
        )
    restored = hs_cfg_now()["sap_url"] == original
    print(
        f"TARGET_RESTORED: system={original_system} -> {'PASS' if restored else 'FAIL'}"
    )
    results.append(restored)

    results.append(case_semi_auto_starts_session(client))

    # Real worklist export + reference produce maintainable groups.
    sample = sorted((ROOT / "downloads").glob("worklist_*.xlsx"))
    if sample:
        from hs_worklist import analyze, read_worklist

        reference = srv.load_reference(hs)
        worklist = read_worklist(sample[-1])
        result = analyze(worklist, reference, hs.get("default_scheme", "EDCHSCDEEX"))
        groups = result["groups"]
        products = sum(g["count"] for g in groups)
        # Every group must carry an HS code and at least one product.
        sane = all(g["hs_code"] and g["count"] > 0 for g in groups)
        ok = len(groups) > 0 and products > 0 and sane
        print(
            f"ANALYSIS: file={sample[-1].name} groups={len(groups)} "
            f"products={products} -> {'PASS' if ok else 'FAIL'}"
        )
        results.append(ok)
    else:
        print("ANALYSIS: SKIP (niciun export in downloads/)")

    r = client.get("/hs/reports")
    report_api_ok = r.status_code == 200 and isinstance(
        (r.get_json() or {}).get("reports"), list
    )
    print(f"REPORT_API: {r.status_code} -> {'PASS' if report_api_ok else 'FAIL'}")
    results.append(report_api_ok)

    srv.HS_REPORT_STORE.start(
        "del-local",
        kind="queue",
        system="FT6",
        items=[{"scheme": "EDCHSCDEEX", "variant": "EU"}],
    )
    remote_delete = client.delete(
        "/hs/reports/del-local", environ_base={"REMOTE_ADDR": "8.8.8.8"}
    )
    local_delete = client.delete("/hs/reports/del-local")
    remote_clear = client.delete("/hs/reports", environ_base={"REMOTE_ADDR": "8.8.8.8"})
    srv.HS_REPORT_STORE.start(
        "del-all",
        kind="queue",
        system="FT6",
        items=[{"scheme": "EDCHSCDEEX", "variant": "EU"}],
    )
    local_clear = client.delete("/hs/reports")
    report_delete_ok = (
        remote_delete.status_code == 403
        and local_delete.status_code == 200
        and remote_clear.status_code == 403
        and local_clear.status_code == 200
        and srv.HS_REPORT_STORE.get("del-local") is None
        and srv.HS_REPORT_STORE.list() == []
    )
    print(
        f"REPORT_DELETE: remote={remote_delete.status_code} "
        f"local={local_delete.status_code} clear_remote={remote_clear.status_code} "
        f"clear_local={local_clear.status_code} -> "
        f"{'PASS' if report_delete_ok else 'FAIL'}"
    )
    results.append(report_delete_ok)

    ok = all(results)
    print("RESULT=", "PASS" if ok else "FAIL")
    srv.HS_REPORT_STORE = original_report_store
    report_tmp.cleanup()
    return 0 if ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
